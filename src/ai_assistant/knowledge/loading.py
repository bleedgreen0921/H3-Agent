"""Local native loaders in a disposable process; explicit loader granularity."""
import asyncio
import json
import os
from pathlib import Path
import tempfile
from langchain_core.documents import Document
from ai_assistant.knowledge import IngestionError
from ai_assistant.knowledge.diagnostics import MAX_BYTES, bounded
from ai_assistant.knowledge.metadata import normalize


def require_content(documents, code, message):
    docs = list(documents)
    valid = [doc for doc in docs if doc.page_content.strip()]
    if not valid:
        raise IngestionError(code, message, f'input_count={len(docs)}; nonempty_count=0')
    return valid


def prepare_documents(loader, *, identity, parser, granularity, splitter, mapping=None, zero_based=()):
    """Accept a native BaseLoader without inventing a parser plugin hierarchy."""
    if granularity not in ('document', 'elements', 'chunks'):
        raise ValueError('Loader granularity must be explicit')
    loaded = require_content(loader.lazy_load(), 'empty_document', '未提取到有效正文。')
    docs = [normalize(d, identity=identity, parser=parser, mapping=mapping, zero_based=zero_based)
            for d in loaded]
    chunks = docs if granularity == 'chunks' else splitter.split_documents(docs)
    chunks = require_content(chunks, 'empty_chunks', '未生成有效片段。')
    for i, chunk in enumerate(chunks):
        chunk.metadata['chunk_index'] = i
    return chunks


MESSAGES = {
    'parser_dependency_error': '本地解析依赖缺失或不兼容，请查看诊断并安装对应依赖。',
    'parser_resource_missing': '本地解析资源缺失，请准备 NLTK、OCR 或其他所需离线资源。',
    'parser_network_blocked': '本地解析器尝试联网，已阻止；请使用完整的离线资源。',
    'parser_timeout': '本地解析超时。',
    'parser_process_failed': '本地解析子进程异常退出，请查看诊断。',
    'parser_protocol_error': '本地解析结果文件无效，请查看诊断。',
    'parser_failed': '本地解析失败，请查看诊断中的具体原因。',
}


class ParserResourceMissing(RuntimeError):
    pass


class ParserNetworkBlocked(RuntimeError):
    pass


def classify_parser_error(error):
    if isinstance(error, ParserResourceMissing) or type(error) is LookupError:
        return 'parser_resource_missing'
    if isinstance(error, ParserNetworkBlocked):
        return 'parser_network_blocked'
    if isinstance(error, ImportError):
        return 'parser_dependency_error'
    return 'parser_failed'


def parser_diagnostic(detail, stderr):
    # Keep the exception tail even when a parser floods stderr.
    detail = bounded(detail, MAX_BYTES // 2)
    prefix = detail + '\nstderr:\n'
    return prefix + bounded(stderr, MAX_BYTES - len(prefix.encode('utf-8')))


async def _stderr_tail(stream):
    tail = bytearray()
    truncated = False
    while data := await stream.read(8192):
        tail.extend(data)
        if len(tail) > MAX_BYTES:
            del tail[:-MAX_BYTES]
            truncated = True
    text = tail.decode('utf-8', errors='replace')
    return bounded(('[stderr 已截断]\n' if truncated else '') + text)


async def _stop_parser(process):
    import signal
    # Kill the group even if its leader exited, so descendants cannot hold stderr open.
    try:
        os.killpg(process.pid, signal.SIGKILL)
    except ProcessLookupError:
        pass
    await process.wait()


async def parse_snapshot(path: Path, *, timeout: float):
    # 独立解释器隔离解析资源，并允许超时终止。
    with tempfile.TemporaryDirectory(prefix='rag-parse-') as directory:
        output = Path(directory) / 'documents.json'
        env = dict(os.environ, HF_HUB_OFFLINE='1',
                   TRANSFORMERS_OFFLINE='1', UNSTRUCTURED_DOWNLOAD_NLTK='False')
        process = await asyncio.create_subprocess_exec(
            __import__('sys').executable, '-m', 'ai_assistant.knowledge.loading', str(path), str(output),
            env=env, stdout=asyncio.subprocess.DEVNULL, stderr=asyncio.subprocess.PIPE,
            start_new_session=True)
        stderr_task = asyncio.create_task(_stderr_tail(process.stderr))
        try:
            try:
                async with asyncio.timeout(timeout):
                    code = await process.wait()
                    stderr = await asyncio.shield(stderr_task)
            except TimeoutError:
                await _stop_parser(process)
                stderr = await stderr_task
                raise IngestionError('parser_timeout', MESSAGES['parser_timeout'], stderr) from None
            if code:
                # Normal child failures have a structured envelope; signals/crashes may not.
                error_code, detail = 'parser_process_failed', f'exit_code={code}\n'
                try:
                    payload = json.loads(output.read_text(encoding='utf-8'))
                    if payload.get('status') == 'error' and payload.get('code') in MESSAGES:
                        error_code = payload['code']
                        detail += str(payload.get('diagnostic', ''))
                except (OSError, ValueError, AttributeError):
                    pass
                raise IngestionError(error_code, MESSAGES[error_code], parser_diagnostic(detail, stderr))
            try:
                payload = json.loads(output.read_text(encoding='utf-8'))
                if payload['status'] != 'ok' or not isinstance(payload['documents'], list):
                    raise ValueError('Invalid parser result envelope')
                return [Document(**item) for item in payload['documents']]
            except (OSError, ValueError, KeyError, TypeError) as exc:
                raise IngestionError('parser_protocol_error', MESSAGES['parser_protocol_error'],
                                     parser_diagnostic(f'{type(exc).__name__}: {exc}', stderr)) from None
        finally:
            await _stop_parser(process)
            # Drain the concurrently consumed pipe before removing temporary files.
            await stderr_task


def _run_local_loader(path):
    # 原生文本/PDF Loader 同样运行在可终止的解析子进程中。
    suffix = Path(path).suffix.lower()
    if suffix in ('.txt', '.md'):
        from langchain_community.document_loaders import TextLoader
        return TextLoader(str(path), encoding='utf-8').load()
    if suffix == '.pdf':
        from langchain_community.document_loaders import PyPDFLoader
        docs = PyPDFLoader(str(path)).load()
        for doc in docs:
            if isinstance(doc.metadata.get('page'), int):
                doc.metadata['page_number'] = doc.metadata['page'] + 1
        return docs
    if suffix == '.csv':
        import csv
        from langchain_core.documents import Document
        with Path(path).open(encoding='utf-8-sig', newline='') as stream:
            rows = csv.reader(stream)
            header = next(rows, [])
            return [Document(page_content=' | '.join(
                f'{header[i] if i<len(header) else i+1}: {value}' for i,value in enumerate(row)),
                metadata={'csv_record':index}) for index,row in enumerate(rows,1) if any(v.strip() for v in row)]
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
        from langchain_core.documents import Document
        book = load_workbook(path, read_only=True, data_only=False)
        values = load_workbook(path, read_only=True, data_only=True)
        docs = []
        try:
            for sheet in book:
                header = None
                for index,(row,display) in enumerate(zip(sheet.iter_rows(values_only=True),
                    values[sheet.title].iter_rows(values_only=True)),1):
                    if header is None:
                        header = row
                    if not any(v is not None for v in row):
                        continue
                    parts = []
                    for i,value in enumerate(row):
                        shown = display[i] if i<len(display) and display[i] is not None else value
                        part = f'{header[i] if i<len(header) else i+1}: {shown}'
                        if isinstance(value,str) and value.startswith('='):
                            part += f'（公式：{value}）'
                        parts.append(part)
                    docs.append(Document(page_content=' | '.join(parts),metadata={
                        'sheet_name':sheet.title,'cell_range':f'A{index}:{get_column_letter(max(1,len(row)))}{index}'}))
            return docs
        finally:
            book.close(); values.close()
    import socket
    def deny_network(*args, **kwargs):
        raise ParserNetworkBlocked('Parser network access is disabled')
    def deny_download(*args, **kwargs):
        raise ParserResourceMissing('Parser resource download is disabled; provision resources offline')
    socket.socket.connect = deny_network
    socket.create_connection = deny_network
    import nltk
    nltk.download = deny_download
    import unstructured.nlp.tokenize as tokenize
    tokenize.download_nltk_packages = deny_download
    from langchain_unstructured import UnstructuredLoader
    loader = UnstructuredLoader(file_path=path, partition_via_api=False,
                                api_key='local-only-unused', url='http://127.0.0.1')
    return list(loader.lazy_load())


def _main():
    import sys
    import traceback
    try:
        docs = _run_local_loader(sys.argv[1])
        payload = dict(status='ok', documents=[
            dict(page_content=d.page_content, metadata=d.metadata) for d in docs])
        code = 0
    except Exception as exc:
        payload = dict(status='error', code=classify_parser_error(exc), diagnostic=bounded(traceback.format_exc()))
        code = 1
    Path(sys.argv[2]).write_text(json.dumps(payload, ensure_ascii=False, default=str), encoding='utf-8')
    raise SystemExit(code)


if __name__ == '__main__':
    _main()
