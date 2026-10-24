"""Real offline parsers, no model, credentials or remote services."""
import asyncio
import pytest


@pytest.mark.parametrize('extension', ['csv', 'docx', 'pptx', 'xlsx', 'md'])
def test_office_formats_parse_offline(tmp_path, extension):
    pytest.importorskip('langchain_unstructured')
    from ai_assistant.knowledge.loading import parse_snapshot
    text = 'Office integration compatibility regression sample.'
    path = tmp_path / ('sample.' + extension)
    if extension == 'csv':
        path.write_text('title,content\nexample,' + text + '\n')
    elif extension == 'md':
        path.write_text('# Sample\n\n' + text)
    elif extension == 'docx':
        from docx import Document
        doc = Document(); doc.add_paragraph(text); doc.save(path)
    elif extension == 'pptx':
        from pptx import Presentation
        doc = Presentation(); slide = doc.slides.add_slide(doc.slide_layouts[1])
        slide.shapes.title.text = 'Sample'; slide.placeholders[1].text = text
        doc.save(path)
    else:
        from openpyxl import Workbook
        doc = Workbook(); doc.active.append(['Title', 'Content']); doc.active.append(['Sample', text]); doc.save(path)
    docs = asyncio.run(parse_snapshot(path, timeout=30))
    assert text in '\n'.join(doc.page_content for doc in docs)
